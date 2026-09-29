from __future__ import annotations

import argparse
from dataclasses import dataclass, asdict
import html
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

META_KEYS = {
    '文書名','システム名','機能id','機能名','機能名称','画面名','画面名称','処理名','api名',
    '版','版数','バージョン','作成者','担当者','作成日','更新日','承認者','文書種別',
    'document name','system name','function id','function name','author','creator','date','version','document type'
}
REQ_TERMS = (
    '必須','以内','以上','以下','禁止','しない','する。','とする','場合','遷移','表示','記録','出力','入力',
    'エラー','ロック','成功','失敗','返却','送信','保存','削除','更新','取得','required','must','shall','when','if ','error'
)
TRIVIAL = {'機能名','作成者','作成日','版数','文書種別','改訂履歴','処理フロー','画面レイアウト','no.','版','日付','内容'}


def norm(s: str) -> str:
    s = html.unescape(str(s or ''))
    s = re.sub(r'<br\s*/?>', ' ', s, flags=re.I)
    s = re.sub(r'<[^>]+>', ' ', s)
    s = s.replace('\\|','|')
    s = unicodedata.normalize('NFKC', s).lower()
    s = re.sub(r'\s+', ' ', s)
    return s.strip()


def clean(v) -> str:
    if v is None:
        return ''
    if hasattr(v, 'isoformat') and not isinstance(v, str):
        try:
            return v.isoformat(sep=' ')
        except TypeError:
            return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace('\r\n','\n').replace('\r','\n').strip()


def meaningful_values(wb) -> list[str]:
    vals = []
    seen = set()
    trivial_norm = {norm(x) for x in TRIVIAL}
    for ws in wb.worksheets:
        nonanchors = set()
        for rng in ws.merged_cells.ranges:
            anchor = str(rng).split(':')[0]
            for row in ws[str(rng)]:
                for cell in row:
                    if cell.coordinate != anchor:
                        nonanchors.add(cell.coordinate)
        for row in ws.iter_rows():
            for c in row:
                if c.coordinate in nonanchors:
                    continue
                t = clean(c.value)
                nt = norm(t)
                if not nt or nt in trivial_norm:
                    continue
                if re.fullmatch(r'[↓↑←→↔┌┐└┘─\s]+', t):
                    continue
                if isinstance(c.value, str) and c.value.startswith('='):
                    continue
                if len(nt) == 1 and not nt.isdigit():
                    continue
                if nt not in seen:
                    seen.add(nt)
                    vals.append(t)
    return vals


def metadata_pairs(wb) -> list[tuple[str, str]]:
    pairs = []
    seen = set()
    for ws in wb.worksheets:
        maxr = min(ws.max_row, 20 if ws.title in {'表紙','カバー'} else 5)
        for r in range(1, maxr + 1):
            items = []
            for c in range(1, ws.max_column + 1):
                v = clean(ws.cell(r, c).value)
                if v:
                    items.append((c, v))
            for i, (_, v) in enumerate(items[:-1]):
                if norm(v) in META_KEYS:
                    nxt = items[i + 1][1]
                    if norm(nxt) not in META_KEYS and nxt:
                        key = (norm(v), norm(nxt))
                        if key not in seen:
                            seen.add(key)
                            pairs.append((v, nxt))
    return pairs


def source_requirements(wb) -> list[str]:
    out = []
    seen = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                t = clean(c.value)
                nt = norm(t)
                if len(t) < 8 or not any(term.lower() in nt for term in REQ_TERMS):
                    continue
                if t.startswith('='):
                    continue
                if nt not in seen:
                    seen.add(nt)
                    out.append(t)
    return out


def sheet_kind(ws, index: int) -> str:
    name = ws.title.lower()
    sample = ' '.join(
        clean(c.value).lower()
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40))
        for c in row if c.value
    )
    if name in {'表紙','カバー','cover'} or (index == 0 and ('仕様書' in sample or 'document' in sample)):
        return 'cover'
    if any(x in name for x in ('フロー','flow')) or '処理フロー' in sample:
        return 'flow'
    if any(x in name for x in ('画面','screen','form')):
        return 'screen'
    if any(x in name for x in ('一覧','list')):
        return 'list'
    if any(x in name for x in ('詳細仕様','詳細設計','機能仕様','article')):
        return 'article'
    return 'grid'


@dataclass
class Evaluation:
    score: float
    passed: bool
    threshold: float
    sheet_coverage: float
    content_coverage: float
    metadata_score: float
    structure_score: float
    requirement_score: float
    ai_readability_score: float
    covered_values: int
    total_values: int
    spreadsheet_artifact_penalty: float
    issues: list[str]
    recommendations: list[str]

    def to_dict(self):
        return asdict(self)


def evaluate(xlsx, markdown, threshold=88.0):
    wb = load_workbook(xlsx, data_only=False, read_only=False)
    md = Path(markdown).read_text(encoding='utf-8')
    nmd = norm(md)
    issues = []
    recs = []

    sheets = [
        ws for ws in wb.worksheets
        if any(c.value not in (None, '') for row in ws.iter_rows() for c in row)
    ]
    sh_hits = 0
    for i, ws in enumerate(sheets):
        kind = sheet_kind(ws, i)
        if norm(ws.title) in nmd:
            sh_hits += 1
        elif kind == 'cover' and re.search(r'(?im)^#\s+\S', md) and re.search(
            r'(文書情報|基本情報|document information|metadata)', md, re.I
        ):
            sh_hits += 1
    sheet_cov = sh_hits / len(sheets) if sheets else 1.0
    if sheet_cov < 1:
        issues.append('One or more source sheets are not traceable in the Markdown.')
        recs.append('Represent each non-empty sheet as a semantic section or source-trace entry.')

    vals = meaningful_values(wb)
    hits = sum(1 for v in vals if norm(v) in nmd)
    content_cov = hits / len(vals) if vals else 1.0
    if content_cov < 0.80:
        issues.append(f'Meaningful source-content coverage is {content_cov:.1%}.')
        recs.append('Recover omitted business facts as normal prose, lists, or domain tables—not as a raw cell dump.')

    pairs = metadata_pairs(wb)
    mhits = sum(1 for _, v in pairs if norm(v) in nmd)
    metadata = mhits / len(pairs) if pairs else 1.0
    if metadata < 0.80:
        issues.append(f'Document/function metadata coverage is {metadata:.1%}.')
        recs.append('Keep function name, author, date, version, IDs, and document metadata in compact semantic blocks.')

    components = [
        1.0 if re.search(r'^#\s+\S', md, re.M) else 0.0,
        1.0 if len(re.findall(r'^##+\s+\S', md, re.M)) >= 3
        else 0.5 if re.search(r'^##\s+\S', md, re.M) else 0.0,
    ]
    kinds = [sheet_kind(ws, i) for i, ws in enumerate(sheets)]
    if 'cover' in kinds:
        components.append(1.0 if re.search(
            r'(文書情報|基本情報|document information|metadata)', md, re.I
        ) else 0.5)
    if 'flow' in kinds:
        mermaid_fence = chr(96) * 3 + 'mermaid'
        components.append(
            1.0 if mermaid_fence in md
            else 0.4 if re.search(r'^\s*\d+[\.)]\s+', md, re.M) else 0.0
        )
    if 'screen' in kinds:
        field_words = sum(1 for w in ('項目名','項目id','必須','桁数','種別') if norm(w) in nmd)
        components.append(
            1.0 if field_words >= 3 and ('|' in md or '<form' in md.lower())
            else 0.4 if field_words >= 2 else 0.0
        )
    if 'list' in kinds:
        components.append(1.0 if re.search(r'^\|.+\|\s*$', md, re.M) else 0.4)
    if 'article' in kinds:
        components.append(1.0 if len(re.findall(r'^###\s+', md, re.M)) >= 3 else 0.5)
    if any('メッセージ' in clean(c.value) for ws in sheets for row in ws.iter_rows() for c in row if c.value):
        components.append(
            1.0 if 'メッセージ' in md and re.search(r'^\|.+\|\s*$', md, re.M) else 0.4
        )
    structure = sum(components) / len(components)
    if structure < 0.75:
        issues.append('The output does not read like a structured software specification.')
        recs.append('Use domain sections such as overview, fields, validations, flow, errors/messages, security, and interfaces.')

    reqs = source_requirements(wb)
    rhits = sum(1 for r in reqs if norm(r) in nmd)
    req_cov = rhits / len(reqs) if reqs else 1.0
    semantic_markers = sum(
        1 for w in ('必須','条件','処理','エラー','セキュリティ','入力','出力','遷移','message','validation')
        if norm(w) in nmd
    )
    explicit = min(1.0, 0.8 * req_cov + 0.2 * min(1.0, semantic_markers / 5))
    if explicit < 0.80:
        issues.append(f'Requirement/rule explicitness is {explicit:.1%}.')
        recs.append('State conditions, constraints, actions, errors, and security rules explicitly; do not rely on visual placement.')

    penalty = 0.0
    bad_patterns = [
        (r'(?im)^###\s+source grid\s*$', 0.30, 'Raw Source grid sections are present.'),
        (r'(?im)^###\s+workbook layout metadata\s*$', 0.25, 'Workbook layout metadata is exposed as specification content.'),
        (r'\bmerge=[A-Z]{1,3}\d+:[A-Z]{1,3}\d+', 0.20, 'Merged-cell coordinates are embedded in the output.'),
        (r'\bcol_width=', 0.20, 'Excel column-width metadata is embedded in the output.'),
        (r'data-grid-width=', 0.15, 'Excel grid-width attributes are embedded in HTML.'),
    ]
    for pat, amount, msg in bad_patterns:
        if re.search(pat, md):
            penalty += amount
            issues.append(msg)

    coord_lines = len(re.findall(r'(?m)^\s*\|\s*\x60?[A-Z]{1,3}\d+\x60?\s*\|', md))
    if coord_lines >= 5:
        penalty += 0.15 if coord_lines < 20 else 0.30
        issues.append(f'{coord_lines} table rows are keyed by Excel cell coordinates.')

    td_count = len(re.findall(r'<t[dh]\b', md, re.I))
    if td_count > 80:
        penalty += 0.15
        issues.append('Large raw HTML grid detected; this resembles spreadsheet layout more than a specification.')

    if any(line.startswith('|') and line.count('|') - 1 > 10 for line in md.splitlines()):
        penalty += 0.10
        issues.append('A very wide table reduces chunkability for AI analysis.')

    if re.search(r'(?im)^###\s+sheet header\s*$', md):
        penalty += 0.10
        issues.append('Sheet headers are rendered as spreadsheet metadata tables instead of document metadata.')

    penalty = min(0.9, penalty)
    readability = max(0.0, 1.0 - penalty)
    if readability < 0.80:
        recs.append('Remove cell coordinates/layout dumps. Preserve meaning with headings, prose, compact domain tables, Mermaid, and useful forms.')

    score = 100 * (
        0.10 * sheet_cov +
        0.20 * content_cov +
        0.10 * metadata +
        0.20 * structure +
        0.20 * explicit +
        0.20 * readability
    )
    score = round(score, 2)
    passed = (
        score >= threshold and
        sheet_cov == 1.0 and
        content_cov >= 0.80 and
        metadata >= 0.80 and
        structure >= 0.75 and
        explicit >= 0.80 and
        readability >= 0.80
    )
    return Evaluation(
        score=score,
        passed=passed,
        threshold=threshold,
        sheet_coverage=round(sheet_cov * 100, 2),
        content_coverage=round(content_cov * 100, 2),
        metadata_score=round(metadata * 100, 2),
        structure_score=round(structure * 100, 2),
        requirement_score=round(explicit * 100, 2),
        ai_readability_score=round(readability * 100, 2),
        covered_values=hits,
        total_values=len(vals),
        spreadsheet_artifact_penalty=round(penalty * 100, 2),
        issues=issues,
        recommendations=recs,
    )


def main():
    p = argparse.ArgumentParser(description='Evaluate XLSX-to-Markdown quality for AI analysis.')
    p.add_argument('xlsx')
    p.add_argument('markdown')
    p.add_argument('--threshold', type=float, default=88.0)
    p.add_argument('--json')
    a = p.parse_args()
    ev = evaluate(a.xlsx, a.markdown, a.threshold)
    data = ev.to_dict()
    print(json.dumps(data, ensure_ascii=False, indent=2))
    if a.json:
        Path(a.json).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    raise SystemExit(0 if ev.passed else 1)


if __name__ == '__main__':
    main()
