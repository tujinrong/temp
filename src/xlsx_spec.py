from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook as _raw_load_workbook
from functools import partial
from .visible_scope import load_visible_workbook, WorkbookScope

# All semantic extraction and evaluation starts from the same visible-only view.
load_workbook = partial(load_visible_workbook, loader=_raw_load_workbook)
from openpyxl.cell.cell import Cell
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


COVER_NAMES = {"表紙", "カバー", "cover", "cover sheet", "front", "front page"}
FLOW_WORDS = ("処理フロー", "処理 flow", "フロー", "flow", "workflow", "process flow", "処理概要")
SCREEN_WORDS = ("画面", "screen", "form", "フォーム", "ui", "レイアウト")
LIST_WORDS = ("一覧", "リスト", "list", "項目", "定義", "table", "テーブル")
ARTICLE_WORDS = ("詳細仕様", "詳細設計", "機能仕様", "説明", "記述", "article", "description")
HEADER_KEYS = (
    "機能名", "機能名称", "function", "function name", "画面名", "画面名称", "帳票名", "api名", "処理名",
    "作成者", "担当者", "author", "creator", "作成日", "更新日", "date", "version", "版", "版数", "文書番号",
)


@dataclass
class CellInfo:
    coordinate: str
    row: int
    col: int
    value: Any
    text: str
    formula: str | None = None
    number_format: str = "General"
    merged_range: str | None = None
    bold: bool = False
    fill: str | None = None
    border: bool = False
    alignment: str | None = None


@dataclass
class SheetInfo:
    name: str
    max_row: int
    max_col: int
    cells: list[CellInfo] = field(default_factory=list)
    merged_ranges: list[str] = field(default_factory=list)
    column_widths: dict[int, float] = field(default_factory=dict)
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[int] = field(default_factory=set)
    charts: int = 0
    images: int = 0
    kind: str = "grid"
    header_rows: int = 0

    @property
    def nonempty(self) -> list[CellInfo]:
        return [c for c in self.cells if c.text != "" or c.formula]


@dataclass
class WorkbookInfo:
    path: Path
    title: str
    sheets: list[SheetInfo]


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time(0, 0, 0) else value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _has_border(cell: Cell) -> bool:
    b = cell.border
    return any(getattr(side, "style", None) for side in (b.left, b.right, b.top, b.bottom))


def _fill_rgb(fill: PatternFill) -> str | None:
    if not fill or fill.fill_type is None:
        return None
    c = fill.fgColor
    if c.type == "rgb" and c.rgb:
        return c.rgb
    if c.indexed is not None:
        return f"indexed:{c.indexed}"
    if c.theme is not None:
        return f"theme:{c.theme}"
    return None


def _merged_lookup(ws) -> dict[str, str]:
    result: dict[str, str] = {}
    for rng in ws.merged_cells.ranges:
        name = str(rng)
        for row in ws[name]:
            for cell in row:
                result[cell.coordinate] = name
    return result


def _contains_term(text: str, term: str) -> bool:
    if re.fullmatch(r"[a-z0-9 _-]+", term, flags=re.I):
        return re.search(r"(?<![a-z0-9])" + re.escape(term) + r"(?![a-z0-9])", text, flags=re.I) is not None
    return term in text


def _detect_kind(name: str, cells: Iterable[CellInfo], index: int) -> str:
    lname = name.strip().lower()
    texts = " ".join(c.text.lower() for c in list(cells)[:80] if c.text)
    if lname in COVER_NAMES or (index == 0 and any(_contains_term(texts, w) for w in ("仕様書", "specification", "システム名", "文書名"))):
        return "cover"
    if any(_contains_term(lname, w) for w in FLOW_WORDS) or any(_contains_term(texts, w) for w in FLOW_WORDS):
        return "flow"
    if any(_contains_term(lname, w) for w in ARTICLE_WORDS):
        return "article"
    if any(_contains_term(lname, w) for w in LIST_WORDS):
        return "list"
    if any(_contains_term(lname, w) for w in SCREEN_WORDS):
        return "screen"
    screen_hits = sum(1 for w in SCREEN_WORDS if _contains_term(texts, w))
    if screen_hits >= 2:
        return "screen"
    return "grid"


def _detect_header_rows(sheet: SheetInfo) -> int:
    hit_rows: list[int] = []
    hit_keys: set[str] = set()
    for c in sheet.nonempty:
        if c.row > 4:
            continue
        raw = c.text.strip().lower().replace("：", ":")
        normalized = re.sub(r"\s+", " ", raw).strip(" :")
        for key in HEADER_KEYS:
            k = key.lower()
            if normalized == k or raw.startswith(k + ":"):
                hit_rows.append(c.row)
                hit_keys.add(k)
                break
    if not hit_rows:
        return 0
    japanese_hit = any(any(ord(ch) > 127 for ch in k) for k in hit_keys)
    if not japanese_hit and len(hit_keys) < 2:
        return 0
    return min(12, max(hit_rows) + 1)


def load_xlsx(path: str | Path) -> WorkbookInfo:
    path = Path(path)
    wb_formula = load_workbook(path, data_only=False, read_only=False)
    try:
        wb_values = load_workbook(path, data_only=True, read_only=False)
    except Exception:
        wb_values = None

    scope = WorkbookScope.from_xlsx(path)
    sheets: list[SheetInfo] = []
    for idx, ws in enumerate(wb_formula.worksheets):
        vws = wb_values[ws.title] if wb_values and ws.title in wb_values.sheetnames else None
        merged = _merged_lookup(ws)
        cells: list[CellInfo] = []
        for row in ws.iter_rows():
            for cell in row:
                formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
                display_value = vws[cell.coordinate].value if (formula and vws) else cell.value
                if formula and display_value is None:
                    display_value = formula
                text = _clean_text(display_value)
                cells.append(
                    CellInfo(
                        coordinate=cell.coordinate,
                        row=cell.row,
                        col=cell.column,
                        value=display_value,
                        text=text,
                        formula=formula,
                        number_format=cell.number_format or "General",
                        merged_range=merged.get(cell.coordinate),
                        bold=bool(cell.font and cell.font.bold),
                        fill=_fill_rgb(cell.fill),
                        border=_has_border(cell),
                        alignment=(cell.alignment.horizontal if cell.alignment else None),
                    )
                )
        widths: dict[int, float] = {}
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            dim = ws.column_dimensions[letter]
            widths[col] = float(dim.width if dim.width is not None else 13.0)
        hidden_cols = {
            col for col in range(1, ws.max_column + 1)
            if not scope.sheets[ws.title].column_visible(col)
        }
        hidden_rows = {row for row in range(1, ws.max_row + 1) if not scope.sheets[ws.title].row_visible(row)}
        info = SheetInfo(
            name=ws.title,
            max_row=ws.max_row,
            max_col=ws.max_column,
            cells=cells,
            merged_ranges=[str(r) for r in ws.merged_cells.ranges],
            column_widths=widths,
            hidden_rows=hidden_rows,
            hidden_cols=hidden_cols,
            charts=len(getattr(ws, "_charts", [])),
            images=len(getattr(ws, "_images", [])),
        )
        info.kind = _detect_kind(info.name, info.nonempty, idx)
        info.header_rows = 0 if info.kind == "cover" else _detect_header_rows(info)
        sheets.append(info)
    return WorkbookInfo(path=path, title=path.stem, sheets=sheets)


def row_map(sheet: SheetInfo, *, start: int = 1, end: int | None = None) -> list[list[CellInfo]]:
    end = end or sheet.max_row
    lookup = {(c.row, c.col): c for c in sheet.cells}
    return [[lookup[(r, c)] for c in range(1, sheet.max_col + 1)] for r in range(start, end + 1)]


def effective_columns(sheet: SheetInfo, narrow_threshold: float = 3.0) -> list[int]:
    """Return columns worth rendering in semantic tables.

    Japanese Excel specifications often use many width=2 columns as a layout grid.
    Narrow empty columns are omitted; narrow columns participating in merges or
    containing data remain represented through their merged anchor or own value.
    """
    used_cols = {c.col for c in sheet.nonempty}
    merged_cols: set[int] = set()
    for rng in sheet.merged_ranges:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", rng)
        if not m:
            continue
        from openpyxl.utils.cell import column_index_from_string
        a = column_index_from_string(m.group(1))
        b = column_index_from_string(m.group(3))
        merged_cols.update(range(a, b + 1))
    result = []
    for col in range(1, sheet.max_col + 1):
        if col in sheet.hidden_cols:
            continue
        width = sheet.column_widths.get(col, 13.0)
        if col in used_cols:
            result.append(col)
        elif width >= narrow_threshold and col not in sheet.hidden_cols:
            result.append(col)
        elif col in merged_cols:
            continue
    return result
