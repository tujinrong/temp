from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = Path("input/japanese_program_spec_test.xlsx")
BLUE, LIGHT, MID, WHITE, GRID = "1F4E78", "D9EAF7", "5B9BD5", "FFFFFF", "D9E2F3"
THIN = Side(style="thin", color="808080")


def merged(ws, rng, value, bold=False, fill=None, center=False):
    ws.merge_cells(rng)
    cell = ws[rng.split(":")[0]]
    cell.value = value
    cell.font = Font(bold=bold, color=WHITE if fill == BLUE else "000000")
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def grid(ws):
    for col in range(1, 27):
        ws.column_dimensions[chr(64 + col)].width = 2
    ws.freeze_panes = "A3"


def header(ws, function_name):
    merged(ws, "A1:D1", "機能名", True, LIGHT)
    merged(ws, "E1:N1", function_name)
    merged(ws, "O1:R1", "作成者", True, LIGHT)
    merged(ws, "S1:Z1", "山田 太郎")
    merged(ws, "A2:D2", "作成日", True, LIGHT)
    merged(ws, "E2:J2", "2026-09-28")
    merged(ws, "K2:N2", "版数", True, LIGHT)
    merged(ws, "O2:R2", "1.0")
    merged(ws, "S2:V2", "文書種別", True, LIGHT)
    merged(ws, "W2:Z2", "プログラム仕様書")


def build(path=OUT):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("表紙")
    for col in range(2, 10):
        ws.column_dimensions[chr(64 + col)].width = 12
    merged(ws, "B3:I5", "販売管理システム\nプログラム仕様書", True, BLUE, True)
    metadata = [
        ("文書名", "ログイン機能 詳細設計・プログラム仕様書"),
        ("システム名", "販売管理システム"),
        ("機能ID", "AUTH-LOGIN-001"),
        ("版数", "1.0"),
        ("作成者", "山田 太郎"),
        ("作成日", "2026-09-28"),
        ("承認者", "佐藤 花子"),
    ]
    for row, (key, value) in enumerate(metadata, 8):
        merged(ws, f"B{row}:D{row}", key, True, LIGHT)
        merged(ws, f"E{row}:I{row}", value)
    merged(ws, "B17:I17", "改訂履歴", True, LIGHT)
    for r, values in enumerate([
        ("版", "日付", "作成者", "内容"),
        ("0.1", "2026-09-20", "山田 太郎", "初版作成"),
        ("1.0", "2026-09-28", "山田 太郎", "承認版"),
    ], 18):
        for c, value in enumerate(values, 2):
            ws.cell(r, c, value)

    ws = wb.create_sheet("機能一覧")
    grid(ws)
    header(ws, "認証機能一覧")
    merged(ws, "A4:Z4", "機能一覧", True, LIGHT)
    spans = [("A", "C"), ("D", "J"), ("K", "N"), ("O", "R"), ("S", "V"), ("W", "Z")]
    for (a, b), value in zip(spans, ["機能ID", "機能名", "種別", "担当", "状態", "備考"]):
        merged(ws, f"{a}6:{b}6", value, True, MID, True)
    rows = [
        ("AUTH-001", "ログイン", "画面", "山田", "完了", "ID/パスワード認証"),
        ("AUTH-002", "ログアウト", "処理", "山田", "完了", "セッション破棄"),
        ("AUTH-003", "パスワード変更", "画面", "佐藤", "作成中", "8文字以上"),
        ("AUTH-004", "アカウントロック解除", "管理", "鈴木", "未着手", "管理者のみ"),
    ]
    for r, values in enumerate(rows, 7):
        for (a, b), value in zip(spans, values):
            merged(ws, f"{a}{r}:{b}{r}", value)
    merged(ws, "A13:H13", "機能数", True, LIGHT)
    merged(ws, "I13:L13", "=COUNTA(A7:A10)")

    ws = wb.create_sheet("画面仕様_ログイン")
    grid(ws)
    header(ws, "ログイン")
    merged(ws, "A4:Z4", "1. 画面レイアウト", True, LIGHT)
    merged(ws, "F7:U7", "ログイン", True, BLUE, True)
    merged(ws, "H10:L11", "ユーザーID", True, LIGHT)
    merged(ws, "M10:T11", "[ user id input ]")
    merged(ws, "H13:L14", "パスワード", True, LIGHT)
    merged(ws, "M13:T14", "[ password input ]")
    merged(ws, "M17:P18", "ログイン", True, GRID, True)
    merged(ws, "Q17:T18", "クリア", True, GRID, True)
    merged(ws, "F21:U22", "エラー時: メッセージを画面上部に表示")
    merged(ws, "A25:Z25", "2. 画面項目定義", True, LIGHT)
    spans2 = [("A", "B"), ("C", "G"), ("H", "J"), ("K", "M"), ("N", "P"), ("Q", "S"), ("T", "V"), ("W", "Z")]
    for (a, b), value in zip(spans2, ["No.", "項目名", "項目ID", "種別", "必須", "桁数", "初期値", "備考"]):
        merged(ws, f"{a}27:{b}27", value, True, MID, True)
    fields = [
        ("1", "ユーザーID", "userId", "text", "○", "50", "", "半角英数字"),
        ("2", "パスワード", "password", "password", "○", "128", "", "マスク表示"),
        ("3", "ログイン", "loginButton", "button", "", "", "", "押下で認証"),
        ("4", "クリア", "clearButton", "button", "", "", "", "入力値を消去"),
    ]
    for r, values in enumerate(fields, 28):
        for (a, b), value in zip(spans2, values):
            merged(ws, f"{a}{r}:{b}{r}", value)

    ws = wb.create_sheet("処理フロー_ログイン")
    grid(ws)
    header(ws, "ログイン認証処理")
    merged(ws, "A4:Z4", "処理フロー", True, LIGHT)
    for rng, value in [
        ("J7:Q8", "開始"),
        ("J11:Q13", "ユーザーID・パスワード\n入力値チェック"),
        ("J16:Q18", "認証API呼出"),
        ("J21:Q23", "認証成功？"),
        ("C27:J29", "エラーメッセージ表示"),
        ("R27:Y29", "メニュー画面へ遷移"),
        ("J33:Q34", "終了"),
    ]:
        merged(ws, rng, value, True, GRID, True)
    for cell, value in [("M9", "↓"), ("M14", "↓"), ("M19", "↓"), ("M24", "┌─ いいえ"), ("Q24", "はい ─┐"), ("M31", "↓")]:
        ws[cell] = value
    merged(ws, "A38:Z38", "補足: 認証失敗が5回連続した場合はアカウントをロックする。")

    ws = wb.create_sheet("詳細仕様_ログイン")
    grid(ws)
    header(ws, "ログイン認証処理")
    sections = [
        (4, "1. 概要", "ユーザーIDとパスワードを用いて利用者を認証し、成功時にメニュー画面へ遷移する。"),
        (9, "2. 入力チェック", "ユーザーIDは必須かつ50文字以内。パスワードは必須かつ128文字以内とする。"),
        (14, "3. 認証処理", "入力チェック成功後、認証APIを呼び出す。APIの戻り値がSUCCESSの場合は認証成功とする。"),
        (20, "4. エラー処理", "認証失敗時はエラーコードに対応するメッセージを画面上部に表示する。"),
        (26, "5. セキュリティ", "パスワードはログへ出力しない。認証失敗回数を記録し、5回連続失敗でアカウントをロックする。"),
    ]
    for row, title, body in sections:
        merged(ws, f"A{row}:Z{row}", title, True, LIGHT)
        merged(ws, f"C{row + 2}:X{row + 4}", body)
    merged(ws, "A34:Z34", "6. メッセージ一覧", True, LIGHT)
    for (a, b), value in zip([("A", "D"), ("E", "L"), ("M", "Z")], ["コード", "条件", "メッセージ"]):
        merged(ws, f"{a}36:{b}36", value, True, MID, True)
    messages = [
        ("E001", "ユーザーID未入力", "ユーザーIDを入力してください。"),
        ("E002", "パスワード未入力", "パスワードを入力してください。"),
        ("E003", "認証失敗", "ユーザーIDまたはパスワードが正しくありません。"),
        ("E004", "アカウントロック", "アカウントがロックされています。管理者へ連絡してください。"),
    ]
    for r, values in enumerate(messages, 37):
        for (a, b), value in zip([("A", "D"), ("E", "L"), ("M", "Z")], values):
            merged(ws, f"{a}{r}:{b}{r}", value)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(path)


if __name__ == "__main__":
    build()
